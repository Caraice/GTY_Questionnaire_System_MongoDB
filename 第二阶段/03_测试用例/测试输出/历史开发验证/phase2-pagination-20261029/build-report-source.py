from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path.cwd();d=Document();sec=d.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    s=d.styles[name];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
    s.element.rPr.rFonts.set(qn('w:eastAsia'),'宋体')
    s.paragraph_format.space_after=Pt(8)
d.styles['Normal'].font.size=Pt(11)
d.styles['Normal'].paragraph_format.line_spacing=1.25
d.styles['Title'].font.size=Pt(21)
d.styles['Heading 1'].font.size=Pt(15)
d.styles['Heading 2'].font.size=Pt(12)
for s in ['Heading 1','Heading 2']:d.styles[s].paragraph_format.keep_with_next=True
sec.header.paragraphs[0].text='MongoDB问卷系统  第二阶段  0.11.0'
f=sec.footer.paragraphs[0];f.alignment=2
f.add_run('第 ');fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld);f.add_run(' 页')
def p(t):d.add_paragraph(t)
def h(t):d.add_heading(t,1)
def table(headers,rows,widths):
    t=d.add_table(rows=1,cols=len(headers));t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for c,s in zip(t.rows[0].cells,headers):c.text=s
    for row in rows:
        for c,s in zip(t.add_row().cells,row):c.text=s
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    t._tbl.tblPr.append(borders)
    for i,row in enumerate(t.rows):
        trp=row._tr.get_or_add_trPr();trp.append(OxmlElement('w:cantSplit'))
        if i==0:trp.append(OxmlElement('w:tblHeader'))
        for c,w in zip(row.cells,widths):
            c.width=Inches(w);c.vertical_alignment=1
            if i==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E9EEF2');c._tc.get_or_add_tcPr().append(sh)
            for pa in c.paragraphs:
                pa.paragraph_format.space_after=Pt(5);pa.paragraph_format.space_before=Pt(5)
                for r in pa.runs:r.font.size=Pt(10);r.bold=i==0
    d.add_paragraph().paragraph_format.space_after=Pt(0)

d.add_paragraph('MongoDB问卷系统第二阶段完成报告','Title')
p('依据最终代码和实际验证整理  2026年10月29日')
p('本阶段围绕题目复用、共享和历史版本，在第一阶段系统上增量修改。八项需求变更已有对应功能和页面入口，最后以四种题型的双账号贯穿测试核对旧问卷、旧答卷和旧统计不变。本文解释设计变化与实际问题；完整接口和逐条用例放在独立文档，避免重复堆砌。')
h('项目目标与原设计的局限')
p('第一阶段把题目嵌入问卷，答案按问卷内题目ID保存。这个方案适合整卷读取、跳转和发布冻结，填写时不用逐题关联查询。但当同一道题要被多个问卷或其他账号复用时，只有问卷内ID不够：无法识别来源，直接修改共享题又可能改变历史答案的含义。')
p('这次我明确保留原分层和MongoDB，要求复用题型校验、保留稳定实例ID、共享不开放答卷，并把涉及用户操作的功能接到页面。具体代码、测试执行和本文整理由AI协助完成。系统没有换成另一个问卷框架，也没有清空旧数据。')
table(['原设计不足','本阶段调整','保留的能力'],[
 ['题目只属于一份问卷','独立题目身份与不可变内容版本','问卷仍嵌入发布快照'],
 ['复制后无法追溯来源','模板ID 版本ID 实例ID分开','原答案键和跳转引用不变'],
 ['修改共享内容可能串改','追加版本 显式升级 按版本共享','已发布问卷继续冻结'],
 ['统计只能按问卷看','按来源版本跨问卷聚合','原整卷 单题统计继续可用']],[1.95,2.55,2.4])
h('系统结构')
p('后端仍是FastAPI和PyMongo，沿用API、Service、Repository分层，前端仍是原生HTML和JavaScript。新增题库、版本历史、共享与引用页面，以及题目跨卷统计入口；原登录、发布、截止、关闭、填写和个人历史继续保留。')

d.add_page_break();h('MongoDB结构与兼容处理')
p('题目身份放question_templates，内容历史放question_versions；题库归属用question_banks和question_bank_memberships分开保存，所以一题能加入多个题库，移出只取消归属。question_grants记录指定版本使用权限，question_usages兼容旧迁移引用，schema_migrations记录迁移。新问卷的question_sources与嵌入题目在同一个文档中原子保存。')
p('我保留嵌入是为了让冻结问卷有完整上下文，同时引用独立版本来管理复用和来源。这里的快照冗余有明确用途，不把所有历史内容塞进单一集合，也不照搬关系数据库的逐题连接。MongoDB适合四题型的不同限制、嵌套规则和问卷快照；集合之间的身份、归属与共享关系仍需要引用。')
p('002迁移采用确定性ID和只插入写法，默认预检，显式apply才写入；重跑不重复造题。原问卷题号、选项ID、答案和HMAC密钥不改写。数字单位是新增可选字段，缺省null不进入规范化指纹；旧版单位不回填，合并时只能明确确认本次单位。')
h('跳转校验与并发恢复')
p('跳转仍是存库的规则：按稳定实例ID定位，首条匹配生效，无匹配顺序继续，支持结束。服务端独立重建路径，不相信客户端路径；升级草稿保留实例ID，再检查选项和跳转引用。单选、多选数量、文本长度、数字范围及整数限制继续复用，拒绝布尔值和非有限数，合法0保留。')
p('内容编号number与管理令牌edit_version分开。保存先用expected_version原子预留完整pending_version，再按固定ID写入不可变版本，最后更新当前指针。两窗口同令牌只有一个胜者，另一方409；版本编号另有唯一索引。单机MongoDB没有跨集合事务，因此恢复补齐同一预留，不通过清锁或重新编号猜测失败状态。')
p('预留、内容或指针写入中断后，重新加载完成验证并恢复；指针已写但回执丢失时，旧令牌重试409。内容碰撞停止恢复，不能覆盖历史。返回本次提交的令牌，避免把较晚写入者的新令牌配给旧内容。数据库连接失败503不等于未落库，页面保留本地内容并要求核对。')

d.add_page_break();h('实际实验与结果')
p('最终在真实MongoDB8.0.17、Python3.14.6上运行全套测试：409 passed，无失败、错误或跳过，131.31秒；Node 25项通过，111.7971毫秒。保留1条既有Starlette/httpx弃用警告。新增13项分页与恢复检查、四题型贯穿4项；这些已包含在409中，不重复相加。')
table(['贯穿步骤','实际结果'],[
 ['A保存并共享v1 B导入发布','四题型都能保存匿名答卷 数字0平均仍为0'],
 ['A创建v2 B回看旧问卷','旧统计JSON相等 旧问卷 答卷 v1 BSON不变'],
 ['B草稿主动升级 A恢复v1','草稿实例ID保留 旧发布卷升级409 恢复生成新v3'],
 ['A撤销v1和v2共享','B新导入均404 空草稿未新增题 已有统计继续可看']],[2.45,4.45])
p('浏览器另用21题库、25版本、同版本22实例和101份文本答卷，实际点击列表、历史、引用和文本翻页。第二页文本仅1条，但总答卷数和回答次数仍为101。跨卷手算样本合并后是3份问卷、3份答卷、5次回答，平均5.4；跳题排除，0计入。')
h('索引测量与缓存判断')
p('固定3000题目、5000版本、2000问卷和10000答卷，在同一合成数据上对比新增索引前后，每项采样7次取中位数。题目深页4.302→1.131毫秒，引用2.503→1.661毫秒；引用扫描量从1000降到100。历史游标只扫描20条，原全量接口读5000版。')
p('跨卷聚合46.160→46.762毫秒，没有证据说明新增索引让统计更快。答卷关联已命中原复合索引，聚合仍有扫描成本。这是本机Repository测量，不含完整HTTP，且有测试并行和热缓存影响，不能当压测结论。因此暂不加额外缓存，避免引入提交、撤权后的失效问题。')

d.add_page_break();h('遇到的问题与修正过程')
p('首先，新增版本不能只写内容后再盲目更新指针，否则单机进程中断会留下半完成状态。我采用完整预留和固定ID补齐，并用写前失败、写后丢回执及两窗口竞争测试验证。随后又检查了响应期间发生较晚写入的情况，让旧内容只返回自己的令牌，下一次过期保存必须409。')
p('其次，列表一次加载全部版本和引用会越来越慢。这轮在Mongo查询中分页，历史改用内容编号游标；来源版本引用增加所有者与版本复合索引。页面测试发现文本翻页后折叠，已保持展开。初轮测试因隔离Mongo未启动出现连接错误，恢复原数据目录后重跑；失败日志保留，没有将其归为通过。')
p('统计中，我把问卷数、答卷数和实例回答次数分开，防止重复导入的题目放大答卷总数。选择题不按文字猜映射，数字题不自动换单位；每份答案先按原版本限制验证。共享题目只授权使用，不能让题目作者顺带看到B的答卷，引用详情也按当前权限过滤。')
h('AI使用与本人核对')
p('AI协助阅读需求、提出结构、编写迁移和业务代码、补充自动测试、执行本地页面检查及整理材料。真实提示、产出、修改和失败过程见AI使用日志。已出现的错误包括统计插入顺序假设、极大数平均处理、页面入口缺失，以及本轮定位折叠项和证据文件路径的问题；具体修正和相应验证按当轮日志保留。')
p('我在对话中明确了沿用第一阶段、稳定ID、权限边界、并发令牌、单位与映射校验、真实记录等要求。本人独立修改过哪些代码、人工审阅了哪些部分，需要我提交前据实补充；本报告不把AI运行的实验写成我已经亲手完成。')
h('教师要求核查与当前边界')
p('两份原文的八项业务方向已有实现和验证。报告已解释第一阶段设计、新问题、结构与程序变化、兼容处理和旧功能回归；详细逐条矩阵见docs/第二阶段/08。跨卷“全部”限定为有结果查看权限的问卷，不开放他人答卷；当前提供原始数字值和平均值，没有额外直方图。')
p('仍未验证全新电脑安装、大规模并发和正式业务恢复；未实现自动单位换算、跨卷人员去重或缓存。第一阶段报告和提交目录保持原样。第二阶段源码、schema、配置与文档、测试证据和本文已具备，但尚未另外整理统一四文件夹提交包；交前还需确认本人分工说明和演示环境。')
p('技术与复现文档：README；docs/第二阶段/01至08。原始本轮证据：baselines/evidence/phase2-pagination-20261029。')
d.core_properties.title='MongoDB问卷系统第二阶段完成报告'
d.core_properties.subject='需求演进 兼容性 并发恢复 分页 贯穿测试'
for tree in [d.element,d.styles.element]:
    for border in list(tree.iter(qn('w:pBdr'))):border.getparent().remove(border)
d.save(root/'deliverables/MongoDB问卷系统第二阶段完成报告.docx')
print('Saved second-stage report; first-stage report unchanged.')
